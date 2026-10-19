"""Manual review (BUILD_SPEC step 5): every reply read by hand; tier and reply verdicts by hand are final.

  export         eval/results.csv -> eval/manual-review.xlsx  (one row per lead, verdict dropdowns)
  import         eval/manual-review.xlsx -> verdicts into eval/results.csv
  export-google  eval/results.csv -> a Google Sheet "Northline manual review" in the runner's Google account
                 (id saved as MANUAL_REVIEW_SHEET_ID in .env; re-export rewrites it)
  import-google  that Google Sheet -> verdicts into eval/results.csv

Import writes manual_tier_ok / manual_reply_ok / manual_notes and prints the final (manual) counts.

Usage:
  .venv/bin/python eval/review_sheet.py export-google
  .venv/bin/python eval/review_sheet.py import-google
"""
import csv
import sys
import urllib.parse
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "eval"))
from envfile import load_env, require, set_env  # noqa: E402
from leads import load_leads  # noqa: E402

RESULTS = ROOT / "eval" / "results.csv"
BOOK = ROOT / "eval" / "manual-review.xlsx"
GOOGLE_TITLE = "Northline manual review"
GOOGLE_TAB = "Review"
COLUMNS = [  # (header, width in characters, source)
    ("lead_id", 8, "lead_id"), ("source", 10, "source"), ("lead message", 46, "_message"),
    ("expected tier", 10, "expected_tier"), ("AI tier", 9, "actual_tier"), ("AI reason", 40, "sheet_reason"),
    ("expected reply", 26, "expected_reply"), ("replied", 9, "replied"), ("seconds to inbox (upper bound)", 12, "seconds_to_inbox"),
    ("auto flags (hints only)", 30, "auto_reply_flags"), ("reply subject", 26, "reply_subject"),
    ("reply text", 70, "reply_text"),
    ("tier correct? (yes/no)", 12, "manual_tier_ok"), ("reply correct? (yes/no/n.a.)", 14, "manual_reply_ok"),
    ("notes", 40, "manual_notes"),
]
SOURCES = [s for _, _, s in COLUMNS]
VERDICT_COLS = ("manual_tier_ok", "manual_reply_ok", "manual_notes")
DROPDOWNS = {"manual_tier_ok": ["yes", "no"], "manual_reply_ok": ["yes", "no", "n.a."]}


def read_results():
    with RESULTS.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def table():
    """Header row + one row per lead, as strings."""
    messages = {l["id"]: l["message"] for l in load_leads()}
    rows = [[messages.get(r["lead_id"], "") if src == "_message" else r.get(src, "") for src in SOURCES]
            for r in read_results()]
    return [h for h, _, _ in COLUMNS], rows


def apply_verdicts(verdicts: dict):
    """verdicts: lead_id -> {manual_tier_ok, manual_reply_ok, manual_notes}. Writes results.csv, prints counts."""
    rows = read_results()
    missing = [r["lead_id"] for r in rows if not verdicts.get(r["lead_id"], {}).get("manual_tier_ok")
               or not verdicts.get(r["lead_id"], {}).get("manual_reply_ok")]
    for r in rows:
        r.update(verdicts.get(r["lead_id"], {}))
    with RESULTS.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    replies = [r for r in rows if r["manual_reply_ok"] in ("yes", "no")]
    print(f"updated {RESULTS.relative_to(ROOT)} with manual verdicts")
    print(f"Tier accuracy (manual): {sum(r['manual_tier_ok'] == 'yes' for r in rows)}/{len(rows)}")
    print(f"Replies judged correct (manual): {sum(r['manual_reply_ok'] == 'yes' for r in replies)}/{len(replies)}"
          f" (n.a.: {sum(r['manual_reply_ok'] == 'n.a.' for r in rows)})")
    for r in rows:
        if r["manual_tier_ok"] == "no" or r["manual_reply_ok"] == "no" or r["manual_notes"]:
            print(f"  {r['lead_id']}: tier {r['manual_tier_ok']}, reply {r['manual_reply_ok']}"
                  f"{' | ' + r['manual_notes'] if r['manual_notes'] else ''}")
    if missing:
        print(f"NOT FINAL: verdicts missing for {', '.join(missing)}")


def clean_verdict(key, value):
    value = str(value or "").strip()
    return value if key == "manual_notes" else value.lower()


# ---------- xlsx ----------

def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    header, rows = table()
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = GOOGLE_TAB
    ws.append(header)
    for row in rows:
        ws.append(row)
    head_fill, verdict_fill = PatternFill("solid", fgColor="1C4A3F"), PatternFill("solid", fgColor="FFF4C7")
    for i, (_, width, src) in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = width
        ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=i).fill = head_fill
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=i)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if src in VERDICT_COLS:
                cell.fill = verdict_fill
    ws.freeze_panes = "B2"
    for src, options in DROPDOWNS.items():
        col = ws.cell(row=1, column=SOURCES.index(src) + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ws.max_row}")
    wb.save(BOOK)
    print(f"wrote {BOOK.relative_to(ROOT)} ({len(rows)} leads). Fill the yellow columns, save, then run: import")


def import_xlsx():
    from openpyxl import load_workbook
    ws = load_workbook(BOOK).active
    assert ws is not None
    pos = {src: [c.value for c in ws[1]].index(h) for h, _, src in COLUMNS}
    apply_verdicts({row[pos["lead_id"]]: {k: clean_verdict(k, row[pos[k]]) for k in VERDICT_COLS}
                    for row in ws.iter_rows(min_row=2, values_only=True)})


# ---------- Google Sheets ----------

def google():
    from observe import Google
    env = load_env()
    require(env, "GOOGLE_CLIENT_ID", "GOOGLE_CLIENT_SECRET", "GOOGLE_RUNNER_REFRESH_TOKEN")
    return env, Google(env["GOOGLE_CLIENT_ID"], env["GOOGLE_CLIENT_SECRET"], env["GOOGLE_RUNNER_REFRESH_TOKEN"])


def export_google():
    from observe import _http
    env, g = google()
    base = "https://sheets.googleapis.com/v4/spreadsheets"
    sheet_id = env.get("MANUAL_REVIEW_SHEET_ID", "")
    if not sheet_id:
        created = _http("POST", base, g._auth(), body={"properties": {"title": GOOGLE_TITLE},
                                                       "sheets": [{"properties": {"title": GOOGLE_TAB}}]}) or {}
        sheet_id = created["spreadsheetId"]
        set_env("MANUAL_REVIEW_SHEET_ID", sheet_id)
    tab_id = next(s["properties"]["sheetId"] for s in
                  (_http("GET", f"{base}/{sheet_id}?fields=sheets.properties", g._auth()) or {})["sheets"]
                  if s["properties"]["title"] == GOOGLE_TAB)
    header, rows = table()
    rng = urllib.parse.quote(f"'{GOOGLE_TAB}'!A1")
    _http("POST", f"{base}/{sheet_id}/values/{urllib.parse.quote(GOOGLE_TAB)}:clear", g._auth(), body={})
    _http("PUT", f"{base}/{sheet_id}/values/{rng}?valueInputOption=RAW", g._auth(), body={"values": [header] + rows})

    n = len(rows) + 1
    cols = len(COLUMNS)
    requests = [
        {"updateSheetProperties": {"properties": {"sheetId": tab_id, "gridProperties": {"frozenRowCount": 1,
                                                                                        "frozenColumnCount": 1}},
                                   "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}},
        {"repeatCell": {"range": {"sheetId": tab_id, "startRowIndex": 0, "endRowIndex": 1, "endColumnIndex": cols},
                        "cell": {"userEnteredFormat": {"backgroundColor": {"red": 0.11, "green": 0.29, "blue": 0.25},
                                                       "textFormat": {"bold": True, "foregroundColor":
                                                                      {"red": 1, "green": 1, "blue": 1}},
                                                       "wrapStrategy": "WRAP", "verticalAlignment": "TOP"}},
                        "fields": "userEnteredFormat"}},
        {"repeatCell": {"range": {"sheetId": tab_id, "startRowIndex": 1, "endRowIndex": n, "endColumnIndex": cols},
                        "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP", "verticalAlignment": "TOP"}},
                        "fields": "userEnteredFormat(wrapStrategy,verticalAlignment)"}},
    ]
    for i, (_, width, src) in enumerate(COLUMNS):
        requests.append({"updateDimensionProperties": {
            "range": {"sheetId": tab_id, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
            "properties": {"pixelSize": width * 7}, "fields": "pixelSize"}})
        if src in VERDICT_COLS:
            requests.append({"repeatCell": {
                "range": {"sheetId": tab_id, "startRowIndex": 1, "endRowIndex": n, "startColumnIndex": i,
                          "endColumnIndex": i + 1},
                "cell": {"userEnteredFormat": {"backgroundColor": {"red": 1, "green": 0.96, "blue": 0.78}}},
                "fields": "userEnteredFormat.backgroundColor"}})
        if src in DROPDOWNS:
            requests.append({"setDataValidation": {
                "range": {"sheetId": tab_id, "startRowIndex": 1, "endRowIndex": n, "startColumnIndex": i,
                          "endColumnIndex": i + 1},
                "rule": {"condition": {"type": "ONE_OF_LIST",
                                       "values": [{"userEnteredValue": v} for v in DROPDOWNS[src]]},
                         "showCustomUi": True, "strict": True}}})
    _http("POST", f"{base}/{sheet_id}:batchUpdate", g._auth(), body={"requests": requests})
    print(f"Google Sheet ready ({len(rows)} leads): https://docs.google.com/spreadsheets/d/{sheet_id}/edit")


def import_google():
    env, g = google()
    require(env, "MANUAL_REVIEW_SHEET_ID")
    values = g.sheet_rows(env["MANUAL_REVIEW_SHEET_ID"], GOOGLE_TAB)
    by_header = {h: src for h, _, src in COLUMNS}
    apply_verdicts({r["lead_id"]: {k: clean_verdict(k, next((v for h, v in r.items() if by_header.get(h) == k), ""))
                                   for k in VERDICT_COLS} for r in values})


if __name__ == "__main__":
    actions = {"export": export_xlsx, "import": import_xlsx, "export-google": export_google,
               "import-google": import_google}
    action = actions.get(sys.argv[1] if len(sys.argv) > 1 else "")
    if not action:
        sys.exit(__doc__)
    action()
